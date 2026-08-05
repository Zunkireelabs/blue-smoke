#!/usr/bin/env python3
"""Generate the Blue Smoke WBS xlsx in the Nepa.works Master Scope format."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HOURLY_RATE = 25  # USD/hr — placeholder; edit before sending to client

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Master Scope"

# ---- styles ----
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
hdr_font = Font(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F3864")
phase_font = Font(bold=True, color="FFFFFF", size=12)
goal_font = Font(italic=True, size=10)
bold = Font(bold=True, size=10)
base_font = Font(size=10)

COLS = ["SN", "Phase", "Module", "Feature Name", "Feature Description (Client-Friendly)",
        "User Roles", "Platform", "Backend Dependency", "Complexity", "Assumptions",
        "Explicit Exclusions", "Risk Note", "Effort (Person-Days)", "Est. Hours",
        "Cost", "Notes / Negotiation Buffer"]
WIDTHS = [5, 10, 14, 30, 50, 14, 12, 14, 11, 34, 28, 34, 12, 10, 12, 30]

row = 1
# title
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
c = ws.cell(row=row, column=1, value="BLUE SMOKE — Vape Device Control App — Master Scope & Work Breakdown")
c.font = Font(bold=True, size=14, color="1F3864")
c.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[row].height = 24
row += 1
ws.cell(row=row, column=1, value=f"Hourly Rate (USD)").font = bold
ws.cell(row=row, column=2, value=HOURLY_RATE).font = bold
ws.cell(row=row, column=4, value="Platform: React Native (iOS+Android) · Verification: on-device ML · Backend: managed BaaS · Firmware: client team (to our spec)").font = goal_font
row += 2

# header row
hdr_row = row
for i, name in enumerate(COLS, start=1):
    cell = ws.cell(row=row, column=i, value=name)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = center
    cell.border = border
ws.row_dimensions[row].height = 30
row += 1

PHASES = [
    {
        "title": "🟦 PHASE 0 — FOUNDATION, ARCHITECTURE & INTEGRATION",
        "fill": "2E75B6",
        "goal": "Establish the app skeleton, managed backend, security/privacy model, and the BLE interface specification (lock/unlock + dead-man auto-lock) that the client firmware team will implement — before any customer-facing feature work begins.",
        "framing": "This phase delivers:\n  •  Full system & data-flow architecture for app + BaaS + device\n  •  A BLE Interface Spec handed to the firmware team\n  •  Managed backend (auth, users, devices, verification flag, push)\n  •  React Native app scaffold + CI/CD for iOS & Android\n  •  Privacy & security design (all ID/biometric data stays on-device)\n  •  Design system and key-screen wireframes",
        "features": [
            ("Architecture", "System Architecture & Data-Flow Design", "Define the full app–BaaS–device architecture: API contracts, data flow, on-device data handling, and integration blueprint for all phases", "Admin", "Backend", "Yes", "High", "Device hardware (YC1012_JD + Cortex-M0+) specs available; BaaS chosen", "Hosting/infra procurement", "Hardware/BLE limitations discovered late", 24, "Architecture doc must be approved before Phase 1 begins"),
            ("BLE", "BLE Interface Spec (Lock/Unlock + Auto-Lock)", "Define the custom GATT service: lock/unlock characteristic, lock-state/battery notify, authenticated session token, and firmware-side dead-man auto-lock on link loss. Delivered as a spec doc for the firmware team", "Admin", "Both", "Partial", "High", "Firmware team can implement the spec on the device SoC", "Firmware development", "Firmware team availability / implementation slippage", 24, "Hard dependency: firmware must implement this before Phase 3 integration"),
            ("Backend", "BaaS Setup — Auth, DB Schema & Push", "Provision managed backend (Supabase/Firebase): authentication, database schema (users, devices, verification flag, audit), and push (FCM/APNs) foundation", "All", "Backend", "Yes", "Medium", "BaaS account available; schema agreed", "Custom backend build", "Vendor limits on auth/push", 16, "Stores only an age_verified flag — never raw ID/biometric data"),
            ("App", "React Native App Scaffold & Native Modules", "Set up the RN app: navigation, TypeScript config, state management, and wiring of native modules (BLE, camera, ML, secure storage)", "All", "Mobile", "No", "Medium", "RN toolchain and native module libraries compatible", "Native (Swift/Kotlin) rewrite", "Native module version conflicts across platforms", 20, "Single codebase for iOS + Android"),
            ("DevOps", "CI/CD & Environments (Dev/Staging/Prod)", "Configure Dev/Staging/Prod environments and CI/CD pipelines with automated builds and signing for both iOS and Android", "Admin", "Both", "No", "Medium", "Apple & Google developer accounts provided by client", "Ongoing infra cost management", "Signing/provisioning issues on iOS", 20, "TestFlight + Play Internal Testing used for demos"),
            ("Security", "Privacy & Security Design", "Design on-device data handling, encrypted secure storage (Keychain/Keystore), HTTPS, token handling, and the data-deletion policy for ID/selfie/biometrics", "Admin", "Both", "Partial", "High", "SSL ready; secure storage APIs available", "Penetration testing (flagged for Phase 3)", "Misconfigured storage or key leakage", 16, "Privacy-by-design is a key client selling point"),
            ("Design", "Design System & Key-Screen Wireframes", "Build the UI kit (components, theme) and wireframe the key flows: onboarding, pairing, verification, lock/unlock", "All", "Mobile", "No", "Medium", "Brand assets provided by client", "Full marketing site design", "Design iterations extend timeline", 24, "Front-loaded so engineering is never blocked on design"),
        ],
        "exit": [
            "- System architecture document approved by all stakeholders",
            "- BLE Interface Spec delivered to and acknowledged by firmware team",
            "- Authentication and database live (users, devices, verification flag)",
            "- React Native app builds on both iOS and Android via CI/CD",
            "- Privacy/security design documented (on-device data, deletion policy)",
            "- Design system and key-screen wireframes approved",
        ],
    },
    {
        "title": "🟩 PHASE 1 — ACCOUNTS & DEVICE MANAGEMENT (BLE)",
        "fill": "548235",
        "goal": "Enable users to create accounts, log in, and pair & manage multiple devices over Bluetooth Low Energy, with a reliable connection lifecycle across foreground and background.",
        "framing": "This phase delivers:\n  •  Account lifecycle: signup, login, password reset\n  •  BLE device scan, discovery, pairing and bonding\n  •  Multiple devices per account (list, rename, unpair, status, battery)\n  •  Device ↔ account sync\n  •  Robust connection lifecycle (reconnect, background)\n  •  Profile & settings",
        "features": [
            ("Auth", "Signup / Login / Password Reset", "Account creation, login, and password reset via the managed backend. Secure session/token handling and logout", "Customer", "Mobile", "Yes", "Medium", "BaaS auth supports email (and optional social)", "SSO / enterprise identity", "Token expiry/refresh edge cases", 20, "Session timeout to be agreed with client"),
            ("Onboarding", "Onboarding & Permissions Flow", "First-run onboarding that explains the product and requests Bluetooth and camera permissions with clear rationale", "Customer", "Mobile", "No", "Low", "Standard OS permission flows", "Animated/video onboarding", "Permission denial handling UX", 12, "Permissions framed to maximise opt-in"),
            ("BLE", "Device Scan & Discovery", "Scan for nearby Blue Smoke devices over BLE and present discoverable devices to the user", "Customer", "Mobile", "No", "Medium", "Device advertises a known service UUID", "Non-BLE transports", "Scanning behaviour differences across OS versions", 18, "Filter to Blue Smoke devices only"),
            ("BLE", "Pairing / Bonding Flow", "Pair and bond with a device, establish the authenticated session, and associate it with the user's account", "Customer", "Mobile", "Partial", "High", "Firmware supports bonding + session token", "Multi-user shared device", "Bonding reliability across iOS/Android", 24, "Device token bound to the verified account"),
            ("Devices", "Multi-Device Management", "Manage multiple paired devices: list, rename, unpair, and view live status (connected/locked, battery)", "Customer", "Mobile", "Yes", "Medium", "Backend stores device records per user", "Device sharing between accounts", "State sync conflicts across devices", 24, "A user can pair several devices"),
            ("Sync", "Device ↔ Account Sync", "Sync paired-device records and state between the app and backend so devices persist across reinstalls/logins", "Customer", "Both", "Yes", "Medium", "Backend device model available", "Cross-user device transfer", "Sync lag / offline edits", 16, "Single source of truth in backend"),
            ("BLE", "Connection Lifecycle & Reconnect", "Handle connect/disconnect, automatic reconnect, and behaviour when the app is backgrounded or the phone sleeps", "Customer", "Mobile", "No", "High", "OS background BLE modes available", "Guaranteed always-on background scanning", "iOS background BLE restrictions", 24, "Foundational for proximity auto-lock in Phase 3"),
            ("Profile", "Profile & Settings", "View account profile and app settings (notifications, paired devices, sign out)", "Customer", "Mobile", "Partial", "Low", "Profile data in backend", "In-depth account editing", "—", 12, "Read-mostly; account edits minimal"),
        ],
        "exit": [
            "- User can sign up, log in, and reset password",
            "- User can scan, pair, and bond with a device over BLE",
            "- User can pair multiple devices and see them listed",
            "- User can rename and unpair devices; live status and battery shown",
            "- Connection reconnects reliably across foreground/background",
            "- Device records persist via backend sync",
        ],
    },
    {
        "title": "🟨 PHASE 2 — AGE & IDENTITY VERIFICATION (ON-DEVICE)",
        "fill": "BF9000",
        "goal": "Gate first-time device activation behind an on-device 18+ check: read date of birth from a government ID and match the ID photo against a live selfie — entirely on-device, with no third-party verification service. This is the highest-risk phase.",
        "framing": "This phase delivers:\n  •  Guided government-ID capture\n  •  On-device OCR to read date of birth and confirm 18+\n  •  Guided selfie capture with basic liveness\n  •  On-device face match: ID photo vs selfie\n  •  Pass/fail decision with retries and a manual fallback policy\n  •  Secure handling and deletion of all images/biometrics\n  •  Verification status that gates first device activation\n\n⚠️  No data leaves the device; only an age_verified flag is stored.",
        "features": [
            ("Capture", "Guided ID Capture UI", "Camera UI with edge detection and guidance so the user can photograph their government ID clearly", "Customer", "Mobile", "No", "Medium", "Device camera adequate; ML edge detection available", "Hardware document scanners", "Glare/low-light reduces capture quality", 18, "Quality of capture directly affects OCR + match accuracy"),
            ("OCR", "ID OCR & DOB Extraction", "On-device text recognition (and PDF417/MRZ barcode parsing where present) to extract the date of birth from the ID", "Customer", "Mobile", "No", "High", "ML Kit (Android) / Vision (iOS) text & barcode APIs", "Government ID validation / authenticity APIs", "Varied ID layouts across regions reduce accuracy", 38, "We verify AGE only; supported ID types to be confirmed per market"),
            ("Age", "Age (18+) Computation & DOB Rules", "Parse extracted DOB across formats, compute age, and enforce the 18+ rule with clear pass/fail", "Customer", "Mobile", "No", "Medium", "DOB reliably extracted by OCR step", "ID expiry / fraud checks", "Date format ambiguity (DD/MM vs MM/DD)", 16, "The only thing verified is age ≥ 18"),
            ("Selfie", "Selfie Capture with Liveness", "Guided, Face-ID-enrollment-style selfie capture with basic liveness (blink/turn/multi-frame) to deter spoofing", "Customer", "Mobile", "No", "High", "Front camera + on-device face detection", "Advanced anti-spoofing (separate add-on)", "Basic liveness can be fooled by sophisticated attacks", 26, "Advanced presentation-attack detection is an upsell"),
            ("Face match", "ID-Photo vs Selfie Face Match", "Extract face embeddings from the ID photo and selfie on-device and compare similarity against a tuned threshold", "Customer", "Mobile", "No", "High", "On-device face embedding available; tuning data accessible", "Cloud face-recognition services", "False rejects on poor ID photos — needs tuning + fallback", 40, "Biggest accuracy risk; dedicated threshold tuning time"),
            ("Flow", "Result Handling, Retries & Fallback", "Manage pass/fail outcomes, retry limits, clear error states, and an agreed manual-review fallback for legitimate edge cases", "Customer", "Mobile", "Partial", "Medium", "Manual fallback policy agreed with client", "Live human review desk (client-operated)", "Fallback policy undefined blocks genuine users", 18, "Manual fallback policy must be agreed with client"),
            ("Privacy", "Secure Handling & Deletion of Biometrics", "Process ID/selfie/embeddings in secure storage and delete them immediately after the pass/fail decision", "Customer", "Mobile", "Partial", "Medium", "Secure storage APIs available", "Long-term document retention", "Residual files if deletion not enforced", 16, "Backend never receives raw images or biometrics"),
            ("Gating", "Verification Status & Activation Gate", "Persist the verification result and gate first-time device activation on a successful 18+ verification", "Customer", "Both", "Yes", "Low", "Backend stores age_verified flag", "Per-session re-verification", "Flag tampering if not server-authoritative", 12, "Only age_verified (+ timestamp/method) is stored"),
        ],
        "exit": [
            "- Valid 18+ ID with a matching selfie → verification passes",
            "- Under-18 ID or mismatched selfie → verification blocked with clear message",
            "- Poor-quality captures → guided retry within retry limits",
            "- All ID/selfie/biometric data deleted after the decision",
            "- Only an age_verified flag stored in the backend",
            "- First-time device activation is gated on successful verification",
        ],
    },
    {
        "title": "🟥 PHASE 3 — PROXIMITY LOCK/UNLOCK & PRODUCTION HARDENING",
        "fill": "C00000",
        "goal": "Deliver secure lock/unlock control and proximity-based auto-lock, validate against real firmware, harden, and launch on both app stores.",
        "framing": "This phase delivers:\n  •  First-time unlock after successful verification\n  •  Authenticated lock/unlock over BLE\n  •  Proximity monitoring + auto-lock when the phone leaves range\n  •  Correct background behaviour on iOS & Android\n  •  Push notifications (lock status, low battery)\n  •  Joint firmware integration testing on real hardware\n  •  Full E2E QA, security review, and store submission",
        "features": [
            ("Activation", "First-Time Unlock (Activation)", "After successful age verification, allow the user to activate and unlock the device for the first time", "Customer", "Mobile", "Partial", "Medium", "Verification flag available; device paired", "Activation without verification", "Activation race conditions", 16, "Activation impossible until age_verified is true"),
            ("Control", "Authenticated Lock/Unlock over BLE", "Send authenticated lock/unlock commands via the BLE characteristic; reflect device-reported lock state in the app", "Customer", "Mobile", "Partial", "High", "Firmware implements the lock command + session auth", "Lock control without a paired phone", "Replay/spoofed commands if auth weak", 24, "Only the bonded, verified user's app can command"),
            ("Proximity", "Proximity Monitor & Auto-Lock on Range Loss", "Monitor BLE connection/RSSI and auto-lock when the phone leaves range; firmware also self-locks on link loss (dead-man)", "Customer", "Both", "Partial", "High", "Firmware dead-man auto-lock implemented per spec", "Precise distance/geofencing", "RSSI is noisy — thresholds need tuning", 38, "Device-enforced lock is the safety guarantee, not app-only"),
            ("Background", "Background BLE Behaviour (iOS & Android)", "Ensure proximity monitoring and auto-lock work when the app is backgrounded or the phone is locked, within OS limits", "Customer", "Mobile", "No", "High", "OS background BLE modes sufficient", "Guaranteed instant background wake on all OEMs", "iOS/Android background BLE restrictions differ", 32, "Behaviour and limits documented for the client"),
            ("Push", "Push Notifications", "Notify the user of relevant events: device locked/unlocked, low battery, out-of-range auto-lock", "Customer", "Both", "Yes", "Medium", "FCM/APNs configured (Phase 0)", "SMS notifications", "Delivery reliability across platforms", 16, "Frequency tuned to avoid spam"),
            ("Integration", "Joint Firmware Integration Testing", "Integrate and test the app against real firmware on a physical device: pairing, lock/unlock, auto-lock, battery, edge cases", "Admin", "Both", "Partial", "High", "Physical device + firmware available by ~week 6", "Firmware bug-fixing (client team)", "Firmware delays compress this window", 24, "Joint testing with the client firmware team"),
            ("QA", "E2E QA, Edge Cases & Security Review", "Full end-to-end testing across the verification, pairing, and lock/unlock flows; security review of auth, storage, and data deletion", "Admin", "Both", "Partial", "High", "Test devices and sample IDs available", "Third-party formal pen-test", "Late-found defects near launch", 24, "Includes the verification accuracy test matrix"),
            ("Release", "App Store Submission (iOS & Android)", "Prepare store listings, assets, privacy disclosures, and submit to the App Store and Google Play", "Admin", "Both", "No", "Medium", "Developer accounts and store assets ready", "Ongoing ASO/marketing", "Store review rejection (age-restricted product policy)", 16, "Age-restricted product policies must be reviewed early", ),
        ],
        "exit": [
            "- Verified user can lock and unlock the device while in BLE range",
            "- Device auto-locks when the phone leaves range (firmware + app)",
            "- Proximity behaviour works when the app is backgrounded",
            "- Push notifications delivered for lock status and low battery",
            "- App validated against real firmware on a physical device",
            "- Passes E2E QA + security review",
            "- Submitted and live on the App Store and Google Play",
        ],
    },
]

ADDONS = {
    "title": "➕ ADD-ONS / UPSELL (OUT OF BASE SCOPE)",
    "fill": "7030A0",
    "goal": "Optional modules presented as add-ons to the core 4-phase delivery.",
    "framing": "Each item below is quoted separately and is not included in the base 2-month scope.",
    "features": [
        ("Admin", "Admin Web Panel", "Web dashboard for device fleet management, verification audit log, and user management", "Admin", "Web", "Yes", "High", "Backend exposes admin APIs", "—", "Standalone product — significant scope", 120, "Recommended for operations at scale"),
        ("Analytics", "Analytics & Crash Reporting", "Integrate analytics and crash reporting (e.g. Firebase/Sentry) for usage and stability insight", "Admin", "Both", "Partial", "Low", "Analytics SDK and account available", "Custom BI dashboards", "—", 24, "Low effort, high operational value"),
        ("Firmware", "Firmware Development", "Develop the device firmware on the YC1012_JD + Cortex-M0+ (BLE service, lock/unlock, dead-man auto-lock)", "Admin", "Both", "No", "High", "Hardware toolchain and access provided", "—", "Embedded scope — separate estimate", 0, "Quoted separately if client wants us to build firmware"),
        ("Security", "Advanced Liveness / Anti-Spoofing", "Presentation-attack detection to harden the selfie step against photo/video/mask spoofing", "Customer", "Mobile", "No", "High", "Suitable on-device PAD model available", "—", "Accuracy/UX tuning intensive", 60, "Strengthens the verification step"),
    ],
}

def write_phase(ws, row, phase, with_exit=True):
    # phase header
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
    c = ws.cell(row=row, column=1, value=phase["title"])
    c.font = phase_font
    c.fill = PatternFill("solid", fgColor=phase["fill"])
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    row += 1
    # goal
    ws.cell(row=row, column=1, value="Goal").font = bold
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=16)
    g = ws.cell(row=row, column=2, value=phase["goal"]); g.font = goal_font; g.alignment = wrap
    ws.row_dimensions[row].height = 48
    row += 1
    # framing
    ws.cell(row=row, column=1, value="Client framing").font = bold
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=16)
    f = ws.cell(row=row, column=2, value=phase["framing"]); f.font = base_font; f.alignment = wrap
    ws.row_dimensions[row].height = 120
    row += 1
    # feature rows
    sn = 1
    first_data = row
    for feat in phase["features"]:
        module, name, desc, roles, plat, dep, cx, assum, excl, risk, hrs, notes = feat
        vals = [sn, phase["title"].split("—")[0].strip()[-2:], module, name, desc, roles, plat, dep, cx,
                assum, excl, risk, f"=N{row}/8", hrs, f"=N{row}*$B$2", notes]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.border = border
            cell.font = base_font
            if i in (5, 10, 11, 12, 16):
                cell.alignment = wrap
            elif i in (1, 6, 7, 8, 9, 13, 14, 15):
                cell.alignment = center
            else:
                cell.alignment = wrap
        ws.row_dimensions[row].height = 78
        sn += 1
        row += 1
    last_data = row - 1
    # subtotal
    ws.cell(row=row, column=4, value="Subtotal").font = bold
    sc = ws.cell(row=row, column=13, value=f"=SUM(M{first_data}:M{last_data})"); sc.font = bold; sc.alignment = center
    hc = ws.cell(row=row, column=14, value=f"=SUM(N{first_data}:N{last_data})"); hc.font = bold; hc.alignment = center
    cc = ws.cell(row=row, column=15, value=f"=SUM(O{first_data}:O{last_data})"); cc.font = bold; cc.alignment = center
    for i in range(1, 17):
        ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor="D9D9D9")
    row += 1
    if with_exit:
        ws.cell(row=row, column=1, value="Exit Criteria").font = bold
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=16)
        e = ws.cell(row=row, column=2, value="\n".join(phase["exit"])); e.font = base_font; e.alignment = wrap
        ws.row_dimensions[row].height = 16 * (len(phase["exit"]) + 1)
        row += 1
    return row, (first_data, last_data)

subtotal_hours_cells = []
row_after = row
for phase in PHASES:
    row_after, _ = write_phase(ws, row_after, phase)
    row_after += 1  # spacer

# Core total
ws.cell(row=row_after, column=4, value="CORE TOTAL (Phases 0–3)").font = Font(bold=True, size=11)
# sum all N cells flagged: easier to sum explicit hours — use a formula across the sheet's N column range
ht = ws.cell(row=row_after, column=14, value=f"=SUMIF(D5:D{row_after-1},\"<>Subtotal\",N5:N{row_after-1})/2")
# simpler & exact: core hours = 668 (sum of feature hours). Use direct values to avoid double counting.
ws.cell(row=row_after, column=13, value=668/8).font = Font(bold=True, size=11)
ws.cell(row=row_after, column=14, value=668).font = Font(bold=True, size=11)
ws.cell(row=row_after, column=15, value=f"=N{row_after}*$B$2").font = Font(bold=True, size=11)
ws.cell(row=row_after, column=14).alignment = center
ws.cell(row=row_after, column=13).alignment = center
ws.cell(row=row_after, column=15).alignment = center
for i in range(1, 17):
    ws.cell(row=row_after, column=i).fill = PatternFill("solid", fgColor="FFE699")
row_after += 1
# buffer line
ws.cell(row=row_after, column=4, value="Negotiation / Risk Buffer (12%)").font = bold
ws.cell(row=row_after, column=14, value=f"=ROUND(N{row_after-1}*0.12,0)").alignment = center
ws.cell(row=row_after, column=15, value=f"=N{row_after}*$B$2").alignment = center
row_after += 1
ws.cell(row=row_after, column=4, value="GRAND TOTAL (Core + Buffer)").font = Font(bold=True, size=11)
ws.cell(row=row_after, column=14, value=f"=N{row_after-2}+N{row_after-1}").font = Font(bold=True, size=11)
ws.cell(row=row_after, column=14).alignment = center
ws.cell(row=row_after, column=15, value=f"=N{row_after}*$B$2").font = Font(bold=True, size=11)
ws.cell(row=row_after, column=15).alignment = center
for i in range(1, 17):
    ws.cell(row=row_after, column=i).fill = PatternFill("solid", fgColor="C6E0B4")
row_after += 2

# Add-ons
row_after, _ = write_phase(ws, row_after, ADDONS, with_exit=False)

# column widths + freeze
for i, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
ws.sheet_view.showGridLines = False

out = "Blue Smoke - WBS - Master Scope.xlsx"
wb.save(out)
print("Saved", out)
